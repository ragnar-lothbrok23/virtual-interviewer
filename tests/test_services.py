from pathlib import Path
from types import SimpleNamespace

import httpx
import numpy as np
import pytest
from docx import Document

from app.services import gemini, resume

SAMPLES = Path(__file__).resolve().parent.parent / "sample_resumes"
# Captured at import time, before conftest's autouse fixture swaps them for a blocker
ORIGINAL_PARSE, ORIGINAL_EMBED = gemini.parse_resume, gemini.embed
ORIGINAL_CLIENT = gemini._client


# ---------- text extraction ----------

def test_extract_text_docx_and_pdf():
    docx_text = resume.extract_text((SAMPLES / "arjun_backend_developer.docx").read_bytes(), "docx")
    pdf_text = resume.extract_text((SAMPLES / "priya_data_analyst.pdf").read_bytes(), "pdf")
    assert "Arjun Mehta" in docx_text and "Flask" in docx_text
    assert "Priya Raman" in pdf_text and "Power BI" in pdf_text


def test_extract_text_reads_docx_tables(tmp_path):
    doc = Document()
    doc.add_paragraph("Candidate Name: Table Person. " * 5)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Skills", "Kubernetes and Terraform"
    path = tmp_path / "t.docx"
    doc.save(path)
    assert "Kubernetes and Terraform" in resume.extract_text(path.read_bytes(), "docx")


def test_extract_text_rejects_near_empty_document(tmp_path):
    doc = Document()
    doc.add_paragraph("Hi")
    path = tmp_path / "short.docx"
    doc.save(path)
    with pytest.raises(resume.ResumeError, match="scanned PDF"):
        resume.extract_text(path.read_bytes(), "docx")


def test_extract_text_rejects_corrupt_and_unknown_files():
    with pytest.raises(resume.ResumeError, match="Could not read"):
        resume.extract_text(b"%PDF-garbage", "pdf")
    with pytest.raises(resume.ResumeError, match="Only PDF and DOCX"):
        resume.extract_text(b"hello", "txt")


# ---------- scoring ----------

@pytest.mark.parametrize("similarity, score", [(0.2, 0.0), (0.60, 0.0), (0.725, 50.0), (0.85, 100.0), (0.99, 100.0)])
def test_similarity_to_score_is_clamped_and_linear(similarity, score):
    assert resume.similarity_to_score(similarity) == score


def test_cosine_similarity():
    a, b = np.array([1.0, 0.0]), np.array([1.0, 1.0])
    assert resume.cosine_similarity(a, a) == pytest.approx(1.0)
    assert resume.cosine_similarity(a, b) == pytest.approx(0.7071, abs=1e-4)


def test_profile_text_handles_empty_sections():
    profile = {"summary": "", "skills": [], "experience": [], "education": [], "projects": [], "certifications": []}
    text = resume.profile_text(profile)
    assert "Skills:" in text and "Projects" not in text


# ---------- Gemini wrapper error handling (no network) ----------

class FakeModels:
    def __init__(self, behaviours):
        self.behaviours = list(behaviours)
        self.calls = 0

    def _next(self):
        self.calls += 1
        b = self.behaviours.pop(0)
        if isinstance(b, Exception):
            raise b
        return b

    def generate_content(self, **kwargs):
        return self._next()

    def embed_content(self, **kwargs):
        return self._next()


@pytest.fixture
def real_gemini_functions(monkeypatch):
    """Undo the conftest block for the wrapper functions themselves; the client underneath is faked."""
    monkeypatch.setattr(gemini, "parse_resume", ORIGINAL_PARSE)
    monkeypatch.setattr(gemini, "embed", ORIGINAL_EMBED)
    monkeypatch.setattr(gemini, "_client", ORIGINAL_CLIENT)  # this test exercises client creation itself
    monkeypatch.setattr(gemini.time, "sleep", lambda s: None)
    return gemini


def use_fake_client(app, module, behaviours):
    models = FakeModels(behaviours)
    app.extensions["gemini_client"] = SimpleNamespace(models=models)
    return models


@pytest.fixture
def job():
    return SimpleNamespace(skill_list=["Python", "SQL"])


def test_embed_retries_dropped_connections(app, real_gemini_functions):
    with app.app_context():
        models = use_fake_client(app, real_gemini_functions, [
            httpx.RemoteProtocolError("Server disconnected"),
            SimpleNamespace(embeddings=[SimpleNamespace(values=[0.1, 0.2])]),
        ])
        vec = real_gemini_functions.embed("text")
        app.extensions.pop("gemini_client")
    assert models.calls == 2 and vec.dtype == np.float32


def test_embed_gives_up_after_three_attempts(app, real_gemini_functions):
    with app.app_context():
        use_fake_client(app, real_gemini_functions, [httpx.ConnectError("down")] * 3)
        with pytest.raises(real_gemini_functions.GeminiError, match="could not"):
            real_gemini_functions.embed("text")
        app.extensions.pop("gemini_client")


def test_quota_error_becomes_friendly_message(app, real_gemini_functions, job):
    from google.genai import errors

    quota = errors.ClientError(429, {"error": {"code": 429, "message": "quota", "status": "RESOURCE_EXHAUSTED"}})
    with app.app_context():
        models = use_fake_client(app, real_gemini_functions, [quota])
        with pytest.raises(real_gemini_functions.GeminiError, match="usage limit"):
            real_gemini_functions.parse_resume("resume text", job)
        app.extensions.pop("gemini_client")
    assert models.calls == 1  # quota errors are not retried


def test_unparseable_response_is_reported(app, real_gemini_functions, job):
    with app.app_context():
        use_fake_client(app, real_gemini_functions, [SimpleNamespace(parsed=None)])
        with pytest.raises(real_gemini_functions.GeminiError, match="unreadable"):
            real_gemini_functions.parse_resume("resume text", job)
        app.extensions.pop("gemini_client")


def test_missing_api_key(app, real_gemini_functions):
    with app.app_context():
        app.extensions.pop("gemini_client", None)
        key = app.config["GEMINI_API_KEY"]
        app.config["GEMINI_API_KEY"] = ""
        try:
            with pytest.raises(real_gemini_functions.GeminiError, match="not configured"):
                real_gemini_functions.embed("x")
        finally:
            app.config["GEMINI_API_KEY"] = key


ORIGINAL_CHECK_CONSISTENCY = gemini.check_consistency


def test_consistency_check_uses_a_fixed_seed_and_default_temperature(app, real_gemini_functions, monkeypatch):
    """Re-checking the same interview must give the same result (seed); temperature stays at Gemini 3's default."""
    monkeypatch.setattr(gemini, "check_consistency", ORIGINAL_CHECK_CONSISTENCY)
    parsed = SimpleNamespace(model_dump=lambda: {"consistency": 9, "summary": "ok", "contradictions": [],
                                                  "resume_mismatches": [], "confidence": "high"})
    with app.app_context():
        models = use_fake_client(app, gemini, [SimpleNamespace(parsed=parsed)])
        seen = []
        original = models.generate_content
        models.generate_content = lambda **kwargs: seen.append(kwargs["config"]) or original(**kwargs)
        result = gemini.check_consistency("Engineer", [("Q1", "Q?", "A."), ("Q2", "Q?", "B.")], {})
        app.extensions.pop("gemini_client")
    assert result["consistency"] == 9
    assert seen[0].seed == app.config["CONSISTENCY_SEED"] and seen[0].temperature is None
