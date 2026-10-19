"""Resume text extraction and resume <-> job matching (paper Section 3.1)."""
import io

import numpy as np
from docx import Document
from flask import current_app
from pypdf import PdfReader

from . import gemini

ALLOWED_EXTENSIONS = {"pdf", "docx"}

# Calibrated on the sample resumes: gemini-embedding-2 gives ~0.67-0.71 for a wrong-role resume
# and ~0.82 for a strong fit, so 0.60-0.85 is stretched linearly onto a 0-100 score.
SIMILARITY_FLOOR = 0.60
SIMILARITY_CEIL = 0.85


class ResumeError(ValueError):
    pass


def extract_text(data, kind):
    """Text from resume bytes. `kind` is 'pdf' or 'docx' (already verified by storage.detect_resume_type)."""
    try:
        if kind == "pdf":
            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
        elif kind == "docx":
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            raise ResumeError("Only PDF and DOCX resumes are supported.")
    except ResumeError:
        raise
    except Exception as e:
        raise ResumeError(f"Could not read the file: {e}") from e

    text = text.strip()
    if len(text) < 100:
        raise ResumeError("Very little text was found. If this is a scanned PDF, please upload a text-based PDF or DOCX.")
    return text


def _project_line(project):
    if isinstance(project, str):  # profiles parsed before Phase 4 stored projects as plain strings
        return project
    parts = [project.get("name", ""), project.get("description", "")]
    if project.get("tech_stack"):
        parts.append(f"Tech: {', '.join(project['tech_stack'])}")
    if project.get("outcome"):
        parts.append(f"Outcome: {project['outcome']}")
    return ". ".join(p for p in parts if p)


def profile_text(profile):
    """The parsed profile flattened into text for embedding (cleaner than raw resume text with addresses etc.)."""
    lines = [f"Summary: {profile.get('summary', '')}", f"Skills: {', '.join(profile.get('skills') or [])}"]
    for exp in profile.get("experience") or []:
        lines.append(f"Experience: {exp['role']} at {exp['organisation']}. {' '.join(exp.get('highlights') or [])}")
    for project in profile.get("projects") or []:
        lines.append(f"Project: {_project_line(project)}")
    for edu in profile.get("education") or []:
        lines.append(f"Education: {edu['degree']}, {edu['institution']}")
    if profile.get("certifications"):
        lines.append(f"Certifications: {'; '.join(profile['certifications'])}")
    return "\n".join(lines)


def skill_summary(profile, required_skills):
    """Coverage and evidence figures shown to HR next to the match score (supporting info, not part of it)."""
    evidence = profile.get("skill_evidence")
    if evidence is None:  # pre-Phase-4 profile: only matched/missing lists exist
        matched = set(profile.get("matched_skills") or [])
        evidence = [{"skill": s, "level": "listed" if s in matched else "missing", "evidence": []} for s in required_skills]
    total = len(evidence) or 1
    proven = sum(e["level"] == "proven" for e in evidence)
    listed = sum(e["level"] == "listed" for e in evidence)
    return {"evidence": evidence, "required": len(evidence), "proven": proven, "listed": listed,
            "missing": len(evidence) - proven - listed,
            "coverage_pct": round(100 * (proven + listed) / total), "proven_pct": round(100 * proven / total)}


def cosine_similarity(a, b):
    return float(np.dot(a, b) / (np.linalg.norm(a) * np.linalg.norm(b)))


def similarity_to_score(similarity):
    scaled = (similarity - SIMILARITY_FLOOR) / (SIMILARITY_CEIL - SIMILARITY_FLOOR)
    return round(float(np.clip(scaled, 0, 1)) * 100, 1)


def job_embedding(job):
    """Job embeddings are cached on the row so each job is embedded only once (cleared when the job is edited)."""
    if job.embedding is None:
        job.embedding = gemini.embed(job.embedding_text).tobytes()
    return np.frombuffer(job.embedding, dtype=np.float32)


def analyse(resume_text, job):
    """Parse the resume with Gemini, then score it against the job.
    Returns (profile, resume_embedding_bytes, similarity, score)."""
    profile = gemini.parse_resume(resume_text, job)
    resume_vec = gemini.embed(profile_text(profile))
    similarity = cosine_similarity(resume_vec, job_embedding(job))
    return profile, np.asarray(resume_vec, dtype=np.float32).tobytes(), similarity, similarity_to_score(similarity)


def score_against(resume_embedding, job):
    """Re-score a stored resume embedding against a (possibly edited) job without calling the parser."""
    similarity = cosine_similarity(np.frombuffer(resume_embedding, dtype=np.float32), job_embedding(job))
    return similarity, similarity_to_score(similarity)
